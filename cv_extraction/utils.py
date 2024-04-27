# cv_extraction/utils.py
import os
import re
import docx
from pdfminer.high_level import extract_text
from openpyxl import Workbook
from openpyxl.utils.exceptions import IllegalCharacterError

def extract_info_from_resume(file_path):
    _, file_extension = os.path.splitext(file_path)
    if file_extension.lower() == '.pdf':
        text = extract_text(file_path)
    elif file_extension.lower() == '.docx':
        doc = docx.Document(file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
    else:
        return None, None, None

    emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    phone_numbers = re.findall(r'(?:(?:\+?(\d{1,3}))?[\s.-]?\(?(?:(\d{3})\)?[\s.-]?)?(\d{3})[\s.-]?(\d{4})(?:\s*(?:ext|x|ext.)\s*(\d{1,5}))?)', text)
    return emails, phone_numbers, text

def process_cvs(cv_directory):
    wb = Workbook()
    ws = wb.active
    ws.append(["Email", "Phone Number", "Text"])

    for filename in os.listdir(cv_directory):
        file_path = os.path.join(cv_directory, filename)
        if os.path.isfile(file_path):
            emails, phone_numbers, text = extract_info_from_resume(file_path)
            if emails:
                for email, phone_number in zip(emails, phone_numbers):
                    try:
                        ws.append([email, " ".join(phone_number), text])
                    except IllegalCharacterError:
                        # Handle IllegalCharacterError by filtering out illegal characters
                        clean_text = ''.join(c for c in text if c.isprintable())
                        ws.append([email, " ".join(phone_number), clean_text])

    wb.save('extracted_info.xlsx')

